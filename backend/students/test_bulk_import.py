import io
import json
import zipfile

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from openpyxl import Workbook, load_workbook
from PIL import Image

from academics.models import AcademicSession, Enrollment, Section
from core.models import AuditLog
from students.bulk_import import COLUMNS
from students.models import Student

pytestmark = pytest.mark.django_db
URL = "/api/v1/students/import/"


def jpeg(colour="blue", size=(1200, 900)):
    buf = io.BytesIO()
    Image.new("RGB", size, colour).save(buf, format="JPEG")
    return buf.getvalue()


def row(**overrides):
    values = {
        "Admission No": "0101", "Admission Date": "01-04-2024", "First Name": "Aarav", "Last Name": "Sharma",
        "Date of Birth": "10-05-2019", "Gender": "Male", "Class": "Class 1", "Section": "A", "Roll No": "1",
        "Address": "12 Gandhi Nagar", "City/Village": "Jaipur", "District": "Jaipur", "PIN Code": "302015",
        "Father Name": "Rohit Sharma", "Father Mobile": "+91 98290 12345", "Mother Name": "Neha Sharma",
        "Mother Mobile": "9829054321", "Primary Contact": "Father", "Parent Aadhaar": "2345 6789 0123",
    }
    values.update(overrides)
    return values


def sheet(*rows, fmt="xlsx"):
    headers = [c.header + (" *" if c.required else "") for c in COLUMNS]
    plain = [c.header for c in COLUMNS]
    if fmt == "csv":
        out = io.StringIO()
        out.write(",".join(headers) + "\n")
        for r in rows:
            out.write(",".join(f'"{r.get(h, "")}"' for h in plain) + "\n")
        return SimpleUploadedFile("students.csv", out.getvalue().encode(), content_type="text/csv")
    book = Workbook()
    ws = book.active
    ws.title = "Students"
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in plain])
    out = io.BytesIO()
    book.save(out)
    return SimpleUploadedFile("students.xlsx", out.getvalue())


def photo_zip(files):
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        for name, content in files.items():
            archive.writestr(name, content)
    return SimpleUploadedFile("photos.zip", out.getvalue(), content_type="application/zip")


@pytest.fixture
def admin(client_for, seeded):
    return client_for("admin")


def test_sample_file_lists_the_schools_classes(admin):
    response = admin.get(URL)
    assert response.status_code == 200
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Instructions", "Students", "Lists"]
    assert book["Students"]["A1"].value == "Admission No"
    classes = [c.value for c in book["Lists"]["A"][1:] if c.value]
    assert classes[0] == "Play Group" and "Class 8" in classes


def test_only_admin_and_principal_can_import(client_for, seeded):
    assert client_for("office").get(URL).status_code == 403
    assert client_for("teacher").post(URL, {"file": sheet(row())}).status_code == 403
    assert client_for("principal").get(URL).status_code == 200


def test_check_reports_every_problem_and_saves_nothing(admin):
    bad = row(**{"Admission No": "0102", "First Name": "", "Date of Birth": "31-02-2019", "Gender": "Boy?",
                 "Class": "Class 12", "Father Mobile": "12345", "Parent Aadhaar": "1234", "PIN Code": "30201"})
    dup = row(**{"Admission No": "0101", "First Name": "Kabir"})
    response = admin.post(URL, {"file": sheet(row(), bad, dup)})
    assert response.status_code == 200
    assert response.data["ok"] is False and response.data["total"] == 3
    problems = {(e["row"], e["column"]) for e in response.data["errors"]}
    assert {(3, "First Name"), (3, "Date of Birth"), (3, "Gender"), (3, "Class"), (3, "Father Mobile"),
            (3, "Parent Aadhaar"), (3, "PIN Code"), (4, "Admission No"), (4, "Roll No")} <= problems
    assert not any(e["row"] == 2 for e in response.data["errors"])
    assert Student.objects.count() == 0


def test_import_creates_students_parents_enrolments_and_photos(admin):
    zipped = photo_zip({"photos/0101.JPG": jpeg(), "anaya.png": jpeg("red"), "extra.jpg": jpeg()})
    second = row(**{"Admission No": "", "First Name": "Anaya", "Gender": "F", "Class": "UKG", "Section": "b",
                    "Roll No": "", "Primary Contact": "Mother", "Photo File": "Anaya.png", "Father Mobile": "9829012345"})
    check = admin.post(URL, {"file": sheet(row(), second), "photos": zipped})
    assert check.data["ok"], check.data["errors"]
    assert check.data["new_sections"] == ["UKG B"]
    assert check.data["photos_matched"] == 2 and check.data["photos_unused"] == ["extra.jpg"]

    zipped.seek(0)
    response = admin.post(URL, {"file": sheet(row(), second), "photos": zipped, "commit": "true"})
    assert response.status_code == 200 and response.data["imported"]
    aarav = Student.objects.get(admission_number="0101")
    anaya = Student.objects.get(first_name="Anaya")
    assert anaya.admission_number == "0001"  # next number from the register
    session = AcademicSession.current()
    assert aarav.enrollments.get().section == Section.objects.get(session=session, school_class__code="1", name="A")
    assert anaya.enrollments.get().section.name == "B"
    father = aarav.guardian_links.get(relationship="father")
    assert father.is_primary_contact and father.guardian.aadhaar_number == "234567890123"
    assert father.guardian.phone == "9829012345"
    # Siblings share the parents' records
    assert anaya.guardian_links.get(relationship="father").guardian == father.guardian
    assert anaya.guardian_links.get(is_primary_contact=True).relationship == "mother"
    with aarav.photo.open("rb") as fh, Image.open(fh) as img:
        assert img.format == "JPEG" and max(img.size) == 600
    assert anaya.photo
    assert AuditLog.objects.filter(action="students.imported").count() == 1
    assert AuditLog.objects.filter(action="student.created").count() == 2



def test_import_is_all_or_nothing(admin):
    response = admin.post(URL, {"file": sheet(row(), row(**{"Admission No": "0102", "Gender": ""})), "commit": "true"})
    assert response.status_code == 400 and not response.data["imported"]
    assert Student.objects.count() == 0


def test_csv_and_excel_number_cells_are_read(admin):
    response = admin.post(URL, {"file": sheet(row(**{"Father Mobile": 9829012345, "Parent Aadhaar": 234567890123.0,
                                                     "PIN Code": 302015, "Roll No": 4}))})
    assert response.data["ok"], response.data["errors"]
    response = admin.post(URL, {"file": sheet(row(), fmt="csv"), "commit": "1"})
    assert response.data["imported"], response.data["errors"]
    assert Enrollment.objects.get().roll_number == 1


def test_existing_admission_number_and_roll_are_refused(admin, school):
    response = admin.post(URL, {"file": sheet(row(**{"Admission No": "S0001", "Roll No": "2"}))})
    columns = {e["column"] for e in response.data["errors"]}
    assert {"Admission No", "Roll No"} <= columns


def test_browser_flow_checks_photo_names_then_uploads_photos_by_admission_number(admin):
    data = {"file": sheet(row(**{"Admission No": "", "Photo File": "aarav"})), "photo_names": json.dumps(["Aarav.jpg"])}
    check = admin.post(URL, data)
    assert check.data["ok"] and check.data["students"][0]["photo"] == "Aarav.jpg"
    missing = admin.post(URL, {"file": sheet(row(**{"Photo File": "nobody.jpg"})), "photo_names": json.dumps(["Aarav.jpg"])})
    assert missing.data["errors"][0]["column"] == "Photo File"

    done = admin.post(URL, {**data, "file": sheet(row(**{"Admission No": "", "Photo File": "aarav"})), "commit": "true"})
    number = done.data["students"][0]["admission_number"]
    upload = SimpleUploadedFile(f"{number}.jpg", jpeg(), content_type="image/jpeg")
    junk = SimpleUploadedFile("9999.jpg", jpeg(), content_type="image/jpeg")
    result = admin.post(f"{URL}photos/", {"photos": [upload, junk]})
    assert result.status_code == 200
    assert [s["admission_number"] for s in result.data["saved"]] == [number]
    assert result.data["unmatched"] == ["9999.jpg"]
    assert Student.objects.get(admission_number=number).photo


def test_photo_upload_rejects_fake_images_and_can_keep_existing(admin, school):
    fake = SimpleUploadedFile("S0001.jpg", b"not really a picture", content_type="image/jpeg")
    result = admin.post(f"{URL}photos/", {"photos": [fake]})
    assert result.data["errors"][0]["file"] == "s0001.jpg"
    admin.post(f"{URL}photos/", {"photos": [photo_zip({"S0001.jpg": jpeg()})]})
    first = Student.objects.get(admission_number="S0001").photo.name
    kept = admin.post(f"{URL}photos/", {"photos": [photo_zip({"S0001.jpg": jpeg("green")})], "replace": "false"})
    assert kept.data["skipped"] == ["s0001.jpg"]
    assert Student.objects.get(admission_number="S0001").photo.name == first


def test_full_section_stops_the_import(admin):
    section = Section.objects.get(session=AcademicSession.current(), school_class__code="1", name="A")
    section.capacity = 1
    section.save()
    rows = [row(), row(**{"Admission No": "0102", "First Name": "Kabir", "Roll No": "2"})]
    response = admin.post(URL, {"file": sheet(*rows), "commit": "true"})
    assert response.status_code == 400 and "holds 1 students" in str(response.data)
    assert Student.objects.count() == 0


def test_command_imports_from_the_server(seeded, tmp_path):
    path = tmp_path / "students.xlsx"
    path.write_bytes(sheet(row()).read())
    photos = tmp_path / "photos"
    photos.mkdir()
    (photos / "0101.jpg").write_bytes(jpeg())
    out = io.StringIO()
    call_command("import_students", str(path), "--photos", str(photos), "--check", stdout=out)
    assert "1 students are ready to import, 1 with photos" in out.getvalue()
    assert Student.objects.count() == 0
    call_command("import_students", str(path), "--photos", str(photos), stdout=out)
    assert Student.objects.get().photo
    call_command("import_students", "--sample", str(tmp_path / "sample.xlsx"), stdout=out)
    assert (tmp_path / "sample.xlsx").exists()


def test_parent_aadhaar_is_masked_for_office_staff(admin, client_for):
    admin.post(URL, {"file": sheet(row()), "commit": "true"})
    student = Student.objects.get()
    detail = client_for("office").get(f"/api/v1/students/{student.pk}/")
    primary = next(g for g in detail.data["guardians"] if g["is_primary_contact"])
    assert primary["guardian"]["aadhaar_number"] == "XXXX-XXXX-0123"
